"""Generate the executive POC one-pager Word document.

Source of truth is markdown_docs/POC_ONE_PAGER.md. Renders a compact,
board-ready .docx (Calibri, accent headings, Light Grid Accent tables) using
the same conventions as generate_system_doc.py.
"""
import re
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SRC = r"C:\sita-platform\markdown_docs\POC_ONE_PAGER.md"
OUT = r"C:\sita-platform\POC_ONE_PAGER.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xB0, 0x1C, 0x1C)
GREEN = RGBColor(0x1A, 0x7F, 0x37)

TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def set_base_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)


def add_runs(para, text):
    """Add text to a paragraph, honouring **bold** and `code` spans."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            para.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            para.add_run(part)


def parse_markdown(md):
    """Split the markdown into (kind, payload) blocks."""
    blocks, lines, i = [], md.splitlines(), 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
            cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
            blocks.append(("table", cells))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
            i += 1
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
            i += 1
        elif line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
            i += 1
        elif line.strip() == "---":
            blocks.append(("hr", None))
            i += 1
        elif line.lstrip().startswith(("- ", "* ")):
            items = []
            while i < len(lines) and lines[i].lstrip().startswith(("- ", "* ")):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).rstrip())
                i += 1
            blocks.append(("ul", items))
        else:
            buf = []
            while i < len(lines) and lines[i].strip() and not re.match(
                r"^\s*(#{1,3}\s|\||[-*]\s|---\s*$)", lines[i]
            ):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("p", " ".join(buf)))
    return blocks


def render(doc, blocks):
    # The paragraph immediately after the first h1 is the tagline, not body copy.
    tagline_at = None
    for idx, (kind, _) in enumerate(blocks):
        if kind == "h1":
            for j in range(idx + 1, len(blocks)):
                if blocks[j][0] in ("p", "ul", "table"):
                    tagline_at = j if blocks[j][0] == "p" else None
                    break
            break

    for idx, (kind, payload) in enumerate(blocks):
        if kind == "h1":
            para = doc.add_heading(payload, 1)
        elif kind == "h2":
            para = doc.add_heading(payload, 2)
        elif kind == "h3":
            para = doc.add_heading(payload, 3)
        elif kind == "p":
            if idx == tagline_at:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(10)
                r = para.add_run(payload.strip("*"))
                r.font.size = Pt(13)
                r.font.color.rgb = GREY
            elif payload.startswith("*") and payload.endswith("*") and not payload.startswith("**"):
                para = doc.add_paragraph()
                r = para.add_run(payload.strip("*"))
                r.italic = True
                r.font.color.rgb = GREY
            else:
                para = doc.add_paragraph()
                add_runs(para, payload)
        elif kind == "ul":
            for item in payload:
                para = doc.add_paragraph(style="List Bullet")
                add_runs(para, item)
        elif kind == "table":
            header, rows = payload[0], payload[1:]
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, hh in enumerate(header):
                cell = t.rows[0].cells[i].paragraphs[0]
                add_runs(cell, hh)
                for r in cell.runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    para = cells[i].paragraphs[0]
                    add_runs(para, v)
                    for r in para.runs:
                        if r.text.strip().startswith(("**Genuinely", "**Genuinely blocked")):
                            r.font.color.rgb = GREEN
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "hr":
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            pPr = para._p.get_or_add_pPr()
            borders = pPr.makeelement(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr", {}
            )
            bottom = borders.makeelement(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bottom",
                {
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val": "single",
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz": "6",
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color": "BFBFBF",
                },
            )
            borders.append(bottom)
            pPr.append(borders)


def main():
    md = open(SRC, encoding="utf-8").read()
    blocks = parse_markdown(md)

    doc = Document()
    set_base_styles(doc)
    sec = doc.sections[0]
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)

    render(doc, blocks)

    foot = doc.sections[0].footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run(
        f"SITA Consolidated Security Reporting · Proof of Concept · "
        f"{date.today().strftime('%d %B %Y')} · Internal"
    )
    r.font.size = Pt(8)
    r.font.color.rgb = GREY

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
