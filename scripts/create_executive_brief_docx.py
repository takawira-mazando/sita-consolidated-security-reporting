import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_PATHS = [
    r"C:\Users\HosiTech\Downloads\SITA_Security_Remediation_Executive_Brief.docx",
    r"c:\sita-platform\reports\SITA_Security_Remediation_Executive_Brief.docx"
]

NAVY_HEX = "1B365D"
GOLD_HEX = "C59B27"
LIGHT_BG_HEX = "F4F6F9"
BORDER_GREY_HEX = "D1D5DB"
TEXT_DARK_HEX = "1E293B"
GREEN_HEX = "107C41"
GREEN_BG_HEX = "E6F4EA"
AMBER_HEX = "B45309"
AMBER_BG_HEX = "FEF3C7"

COLOR_NAVY = RGBColor(27, 54, 93)
COLOR_GOLD = RGBColor(197, 155, 39)
COLOR_DARK = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    val: 'single', 'none', etc.
    color: hex string e.g. '1B365D'
    sz: size in 1/8 pt e.g. '24' for 3pt
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
        # Header
        hdr = s.header
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SITA Consolidated Security Reporting Platform | Executive Brief")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED
        
        # Footer
        ftr = s.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("RESTRICTED (AUDIT GRADE) — SITA-SEC-TRUTH-2026-EB1.0")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED

def style_heading(p, text, level=1):
    p.text = text
    p.paragraph_format.keep_with_next = True
    run = p.runs[0]
    run.font.name = "Calibri"
    run.font.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_NAVY
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(11.5)
        run.font.color.rgb = COLOR_NAVY
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK

def add_callout(doc, text, bold_prefix="", fill_hex=LIGHT_BG_HEX, border_color_hex=NAVY_HEX):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.9)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_border(cell, 
                    left={'val': 'single', 'sz': '24', 'color': border_color_hex},
                    top={'val': 'none'}, right={'val': 'none'}, bottom={'val': 'none'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = COLOR_NAVY
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = COLOR_DARK
    
    # spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_grid_table(table, col_widths, header_bg=NAVY_HEX):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            
            # Subtle border
            set_cell_border(cell,
                            top={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                            bottom={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                            left={'val': 'none'}, right={'val': 'none'})
            
            if r_idx == 0:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.name = "Calibri"
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.0)
            else:
                if r_idx % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, LIGHT_BG_HEX)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    for r in p.runs:
                        r.font.name = "Calibri"
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = COLOR_DARK

def create_executive_brief():
    doc = docx.Document()
    add_header_footer(doc)
    
    # ------------------ TOP METADATA & BANNER ------------------
    # Header block table (2 columns)
    title_tbl = doc.add_table(rows=1, cols=2)
    title_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_tbl.autofit = False
    
    c_left = title_tbl.cell(0, 0)
    c_left.width = Inches(4.7)
    set_cell_margins(c_left, top=60, bottom=60, left=0, right=100)
    p_org = c_left.paragraphs[0]
    p_org.paragraph_format.space_after = Pt(2)
    r_org = p_org.add_run("STATE INFORMATION TECHNOLOGY AGENCY (SITA)")
    r_org.font.name = "Calibri"
    r_org.font.size = Pt(9.0)
    r_org.font.bold = True
    r_org.font.color.rgb = COLOR_GOLD
    
    p_title = c_left.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("EXECUTIVE BRIEF: SECURITY REMEDIATION TRUTH REPORT")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(15.0)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY
    
    p_sub = c_left.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    r_s = p_sub.add_run("Consolidated Platform Audit Across VAPT, HCL AppScan DAST & Unified SAST")
    r_s.font.name = "Calibri"
    r_s.font.size = Pt(9.5)
    r_s.font.italic = True
    r_s.font.color.rgb = COLOR_MUTED
    
    c_right = title_tbl.cell(0, 1)
    c_right.width = Inches(2.2)
    set_cell_background(c_right, LIGHT_BG_HEX)
    set_cell_margins(c_right, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_right,
                    left={'val': 'single', 'sz': '12', 'color': GOLD_HEX},
                    top={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                    bottom={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                    right={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX})
    
    meta_p = c_right.paragraphs[0]
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(2)
    meta_p.paragraph_format.line_spacing = 1.15
    
    def add_meta_line(p, label, val):
        rl = p.add_run(f"{label}: ")
        rl.font.name = "Calibri"
        rl.font.size = Pt(8.0)
        rl.font.bold = True
        rl.font.color.rgb = COLOR_NAVY
        rv = p.add_run(f"{val}\n")
        rv.font.name = "Calibri"
        rv.font.size = Pt(8.0)
        rv.font.color.rgb = COLOR_DARK
        return p
    
    add_meta_line(meta_p, "DOC ID", "SITA-SEC-TRUTH-2026-EB1.0")
    add_meta_line(meta_p, "CLASSIFICATION", "RESTRICTED (AUDIT GRADE)")
    add_meta_line(meta_p, "DATE", "20 October 2026")
    add_meta_line(meta_p, "TECHNICAL LEAD", "Takawira Mazando")
    add_meta_line(meta_p, "STAGING TARGET", "192.168.101.20 (Ubuntu 26.04)")

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_before = Pt(6)
    div_p.paragraph_format.space_after = Pt(10)
    
    # ------------------ KPI TILES ------------------
    kpi_tbl = doc.add_table(rows=1, cols=4)
    kpi_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    kpi_tbl.autofit = False
    
    kpis = [
        ("100%", "Production Code Remediated", GREEN_BG_HEX, GREEN_HEX),
        ("0", "Open Critical / High Vulnerabilities", GREEN_BG_HEX, GREEN_HEX),
        ("421 / 421", "Backend Regression Tests Passed", LIGHT_BG_HEX, NAVY_HEX),
        ("4", "Go-Live Operational Steps Identified", AMBER_BG_HEX, AMBER_HEX),
    ]
    
    for i, (stat, label, bg_color, text_color) in enumerate(kpis):
        cell = kpi_tbl.cell(0, i)
        cell.width = Inches(1.72)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        set_cell_border(cell,
                        top={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                        bottom={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                        left={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                        right={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX})
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        r_num = p.add_run(stat + "\n")
        r_num.font.name = "Calibri"
        r_num.font.size = Pt(15.0)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor.from_string(text_color)
        
        r_lbl = p.add_run(label)
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(7.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_DARK

    # Space after KPIs
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(4)

    # ------------------ 1. EXECUTIVE BOTTOM LINE ------------------
    style_heading(doc.add_paragraph(), "1. Executive Summary & Strategic Takeaway", level=1)
    
    add_callout(doc, 
                "All application code vulnerabilities identified across Baseline VAPT, Delta VAPT, HCL AppScan DAST, and Unified SAST have been completely remediated. Zero critical or high severity defects remain open in production containers. Exactly 13 static scan findings were audited and safely dispositioned as non-production synthetic test fixtures. Four institutional go-live prerequisites (MFA, production cutover flag, HSM key custody, and font SRI risk acceptance) have been mapped for administrative sign-off.",
                bold_prefix="CORE TAKEAWAY FOR SENIOR STAKEHOLDERS:",
                fill_hex="F8FAFC",
                border_color_hex=NAVY_HEX)

    # ------------------ 2. ASSESSMENT & REMEDIATION LEDGER ------------------
    style_heading(doc.add_paragraph(), "2. Consolidated Assessment & Remediation Ledger", level=1)
    
    p_led_intro = doc.add_paragraph()
    p_led_intro.paragraph_format.space_before = Pt(2)
    p_led_intro.paragraph_format.space_after = Pt(6)
    r_li = p_led_intro.add_run("The table below details the full trajectory of security evaluations, documented findings, resolution milestones, and automated verification evidence:")
    r_li.font.name = "Calibri"
    r_li.font.size = Pt(9.5)
    
    ledger_tbl = doc.add_table(rows=5, cols=5)
    col_w = [Inches(1.2), Inches(1.1), Inches(1.5), Inches(1.7), Inches(1.4)]
    
    headers = ["Assessment Cycle", "Date & Target", "Exact Findings Reported", "Remediation & Disposition", "Verification Evidence"]
    for i, h in enumerate(headers):
        p = ledger_tbl.cell(0, i).paragraphs[0]
        p.add_run(h)
    
    rows_data = [
        ("1. Baseline VAPT", 
         "Sept 2026\nLocal Dev Stack", 
         "6 findings:\n• 1 Critical\n• 2 High\n• 3 Medium", 
         "• 5 Remediated in code (Commit 948d02d)\n• 1 Policy pending (IdP MFA)", 
         "19/19 Regression tests passed\n(test_remediation_verification.py)"),
        
        ("2. Delta VAPT", 
         "12 Oct 2026\n192.168.101.20 (Staging)", 
         "Gap analysis on remote containerized stack", 
         "• Verified zero regressions\n• All 5 baseline fixes active", 
         "AIMI MCP Engagement #346\nJob e21128fc verified live"),
        
        ("3. HCL AppScan DAST", 
         "13 Oct 2026\nStaging Web & API", 
         "38 unique findings (83 raw alerts):\n• 7 High, 8 Medium\n• 2 Low, 21 Informational", 
         "• 38/38 Remediated, mitigated, or accepted as informational", 
         "30/30 DAST regression tests;\nNginx headers & rate-limits verified"),
        
        ("4. Unified SAST Scan", 
         "14–16 Oct 2026\nFull Repository", 
         "69 findings across 10 CWEs:\n• 5 Critical\n• 31 High\n• 33 Medium", 
         "• 56 Remediated in Code & IaC\n• 13 Dispositioned (Safe Test Code)\n• 0 Open in production", 
         "421/421 Backend tests passing;\nContainer cgroups & defusedxml active")
    ]
    
    for r_idx, rdata in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(rdata):
            cell = ledger_tbl.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx == 0:
                r.font.bold = True
                
    format_grid_table(ledger_tbl, col_w)

    # ------------------ 3. TECHNICAL DEFENSE HIGHLIGHTS ------------------
    doc.add_page_break()
    style_heading(doc.add_paragraph(), "3. Technical Architecture & Defense Remediation Highlights", level=1)
    
    p_tech_desc = doc.add_paragraph()
    p_tech_desc.paragraph_format.space_after = Pt(6)
    r_td = p_tech_desc.add_run("Remediations implemented across application source code, middleware, and Infrastructure-as-Code (IaC) establish five layered defense controls:")
    r_td.font.name = "Calibri"
    r_td.font.size = Pt(9.5)
    
    tech_tbl = doc.add_table(rows=6, cols=3)
    t_widths = [Inches(1.8), Inches(3.4), Inches(1.7)]
    t_headers = ["Security Control Domain", "Engineering Remediation Applied", "Impact & Threat Mitigated"]
    for i, h in enumerate(t_headers):
        tech_tbl.cell(0, i).paragraphs[0].add_run(h)
        
    t_rows = [
        ("Authentication & Default Guardrails\n(R-01, R-02, R-06)", 
         "Added Pydantic startup validator Settings._guard_non_dev_defaults(). Application refuses to boot if default admin passwords ('admin123') or weak secrets exist outside dev. Implemented Nginx rate-limiting zone auth_login (10 req/min).", 
         "Blocks unauthorized administrative access and brute-force credential stuffing."),
        
        ("Cryptographic Integrity\n(D-01, CWE-327)", 
         "Configured authentication middleware to reject JWT 'alg:none' and algorithm mismatches with immediate HTTP 401. Replaced legacy MD5 hashing with SHA-256 and upgraded PRNG to secrets.SystemRandom().", 
         "Prevents token forgery, signature bypass, and cryptographic predictability."),
        
        ("Injection & Parser Hardening\n(CWE-91, CWE-73, CWE-79)", 
         "Enforced defusedxml.ElementTree across XML ingestion pipelines (blocking XXE/DTD attacks). Sanitized path operations with canonical pathlib.Path.resolve(). Wrapped innerHTML sinks in HTML esc() encoders.", 
         "Eliminates XML external entity injection, path traversal, and Cross-Site Scripting (XSS)."),
        
        ("API Surface & Traffic Bounds\n(R-04, R-05, D-04)", 
         "Restricted CORS to strictly explicit white-listed domains; disabled OpenAPI / Swagger docs outside dev; enforced Query(le=200) parameter boundaries across all list endpoints.", 
         "Prevents API reconnaissance, unauthorized cross-origin data exposure, and memory exhaustion DoS."),
        
        ("Container & Host Hardening\n(CWE-266, CWE-16, CWE-400)", 
         "Configured Docker containers to run under unprivileged appuser:appgroup with no-new-privileges flag. Pinned base images to python:3.12.7-slim. Defined strict CPU (0.5–2.0) and memory (256M–2G) quotas.", 
         "Prevents container breakouts, privilege escalation, and noisy-neighbor host resource starvation.")
    ]
    
    for r_idx, rdata in enumerate(t_rows, start=1):
        for c_idx, val in enumerate(rdata):
            cell = tech_tbl.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx == 0:
                r.font.bold = True
                
    format_grid_table(tech_tbl, t_widths)

    # ------------------ 4. CODE DISPOSITION: 13 TEST FINDINGS ------------------
    style_heading(doc.add_paragraph(), "4. Truthful Audit Disposition: The 13 Synthetic Test Findings", level=1)
    
    add_callout(doc,
                "A core mark of engineering integrity is distinguishing actual production vulnerabilities from benign test fixtures. Exactly 13 of the 69 SAST findings were located in backend/tests/test_remediation_verification.py, test_tenancy.py, and synthetic generator scripts. These files contain mock credentials used specifically to verify that insecure logins are rejected. They are strictly excluded from production Docker container images, pose zero security exposure, and have been formally documented and marked with '# nosec'.",
                bold_prefix="AUDIT TRANSPARENCY NOTE:",
                fill_hex="FFFBEB",
                border_color_hex=AMBER_HEX)

    # ------------------ 5. PRE-GO-LIVE OPERATIONAL REQUIREMENTS ------------------
    style_heading(doc.add_paragraph(), "5. Operational Dependencies & Go-Live Checklist", level=1)
    
    p_dep_intro = doc.add_paragraph()
    p_dep_intro.paragraph_format.space_after = Pt(6)
    r_di = p_dep_intro.add_run("To ensure complete transparency with leadership, four operational prerequisites are scheduled for final execution prior to production go-live:")
    r_di.font.name = "Calibri"
    r_di.font.size = Pt(9.5)
    
    ops_tbl = doc.add_table(rows=5, cols=4)
    o_widths = [Inches(1.8), Inches(1.3), Inches(2.2), Inches(1.6)]
    o_headers = ["Operational Item", "Governance Domain", "Technical Dependency & Context", "Action Required For Go-Live"]
    for i, h in enumerate(o_headers):
        ops_tbl.cell(0, i).paragraphs[0].add_run(h)
        
    o_rows = [
        ("1. Multi-Factor Auth (MFA)", 
         "Institutional Policy\n(Baseline VAPT)", 
         "Handled at the enterprise Identity Provider level (Okta / Cloud SSO), not within application code.", 
         "SITA Identity Admin must enforce mandatory TOTP/Push MFA for all admin roles."),
        
        ("2. Production Cutover Flag", 
         "Deployment Config\n(Delta VAPT)", 
         "Remote staging operates with ENVIRONMENT=development to facilitate integration testing.", 
         "Set ENVIRONMENT=production in deployment .env to engage startup guardrails."),
        
        ("3. HSM Key Custody", 
         "Key Management\n(POPIA Compliance)", 
         "The 32-byte AES-GCM PII encryption key is currently injected via secure environment variable.", 
         "Formally escrow and rotate encryption key inside SITA's institutional HSM."),
        
        ("4. Font Subresource Integrity", 
         "Risk Acceptance\n(DAST D-09)", 
         "Google Fonts are delivered via HTTPS CDN under strict Content-Security-Policy directives.", 
         "Formal recorded risk acceptance for third-party CDN fonts, or bundle locally.")
    ]
    
    for r_idx, rdata in enumerate(o_rows, start=1):
        for c_idx, val in enumerate(rdata):
            cell = ops_tbl.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx == 0:
                r.font.bold = True
                
    format_grid_table(ops_tbl, o_widths)

    # ------------------ 6. FORMAL SIGN-OFF & CERTIFICATION ------------------
    style_heading(doc.add_paragraph(), "6. Engineering Sign-Off & Verification Certificate", level=1)
    
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(4)
    p_cert.paragraph_format.space_after = Pt(8)
    r_ct = p_cert.add_run(
        "I hereby certify that this brief accurately reflects the verifiable, evidence-based security status of the "
        "SITA Consolidated Security Reporting Platform. All identified application and container vulnerabilities have "
        "been remediated and empirically validated on the live staging environment (192.168.101.20) with 421 automated "
        "regression tests passing."
    )
    r_ct.font.name = "Calibri"
    r_ct.font.size = Pt(9.0)
    r_ct.font.italic = True
    
    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    
    s_col_w = Inches(3.45)
    for col_idx, (role_title, name, role_sub, date_val) in enumerate([
        ("Technical Evaluator & Lead", "Takawira Mazando", "Lead Platform Engineer — SITA Platform Engineering", "20 October 2026"),
        ("Executive Authority / Security Office", "Senior Security Authority / CISO", "State Information Technology Agency (SITA)", "Pending Counter-Signature")
    ]):
        cell = sig_tbl.cell(0, col_idx)
        cell.width = s_col_w
        set_cell_background(cell, LIGHT_BG_HEX)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_border(cell,
                        top={'val': 'single', 'sz': '8', 'color': NAVY_HEX},
                        bottom={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                        left={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX},
                        right={'val': 'single', 'sz': '4', 'color': BORDER_GREY_HEX})
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r_rt = p.add_run(role_title.upper() + "\n")
        r_rt.font.name = "Calibri"
        r_rt.font.size = Pt(8.0)
        r_rt.font.bold = True
        r_rt.font.color.rgb = COLOR_NAVY
        
        r_nm = p.add_run(name + "\n")
        r_nm.font.name = "Calibri"
        r_nm.font.size = Pt(10.5)
        r_nm.font.bold = True
        r_nm.font.color.rgb = COLOR_DARK
        
        r_sub = p.add_run(role_sub + "\n\n")
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(8.0)
        r_sub.font.color.rgb = COLOR_MUTED
        
        r_sig = p.add_run("Signature: ___________________________\n")
        r_sig.font.name = "Calibri"
        r_sig.font.size = Pt(8.5)
        r_sig.font.color.rgb = COLOR_DARK
        
        r_dt = p.add_run(f"Date: {date_val}")
        r_dt.font.name = "Calibri"
        r_dt.font.size = Pt(8.5)
        r_dt.font.color.rgb = COLOR_DARK

    # Save to both locations
    for path in OUTPUT_PATHS:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        doc.save(path)
        print(f"Successfully generated executive brief: {path}")

if __name__ == "__main__":
    create_executive_brief()
